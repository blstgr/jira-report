#!/usr/bin/env python3
import datetime as dt
import json
import os
import re
import shlex
import subprocess
import sys
import threading
import time
from pathlib import Path
import urllib.error
import urllib.request

from google_drive_sync import authorize, is_placeholder_client_config, upload_or_update

try:
    import readline

    readline.parse_and_bind("set editing-mode emacs")
    readline.parse_and_bind("tab: complete")
except Exception:
    readline = None


APP_DIR = Path(__file__).resolve().parent
ROOT = APP_DIR.parent
SETTINGS_DIR = ROOT / "settings"
SETTINGS_TEMPLATE = SETTINGS_DIR / "roadmap-settings.json"
STATE_PATH = SETTINGS_DIR / "roadmap-settings.local.json"
REPORTS_DIR = ROOT / "report"
OAUTH_SETUP_FILE = APP_DIR / "google-oauth-setup.json"
JIRA_SERVICE = os.environ.get("JIRA_KEYCHAIN_SERVICE", "atlassian-dc-mcp")
JIRA_ACCOUNT = os.environ.get("JIRA_KEYCHAIN_ACCOUNT", "jira-token")
JOB_NAME = "roadmap-jira-report-update"
GOOGLE_CLIENT_SECRETS_CANDIDATES = [
    APP_DIR / "google-oauth-client-secrets.json",
    APP_DIR / "google-client-secrets.json",
    APP_DIR / "credentials.json",
    ROOT / "google-oauth-client-secrets.json",
    ROOT / "google-client-secrets.json",
    ROOT / "credentials.json",
]
SYMBOLS = ["◐", "◓", "◑", "◒", "✦", "✧", "⬣", "⬢"]
TERM_WIDTH = 100
DEFAULT_DONE_STATUSES = ["Done", "QA Prod Done", "In Validation"]
BACK_COMMANDS = {"/back", "/b"}
EDIT_COMMANDS = {"/edit", "/e"}
CLEAR_COMMANDS = {"/clear", "/c"}
MALFORMED_EDIT_PREFIXES = ("e/", "edit/")
URL_PATTERN = re.compile(r"(https?://[^\s]+)")


def terminal_clickable(text):
    def repl(match):
        url = match.group(1)
        return f"\033]8;;{url}\033\\{url}\033]8;;\033\\"
    return URL_PATTERN.sub(repl, text)


def normalize_keyword(text):
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def spinner_line(symbol, message):
    text = f"{symbol} {message}"
    if len(text) > TERM_WIDTH - 1:
        text = text[: TERM_WIDTH - 4] + "..."
    return f"\r\x1b[2K{text}"


def status_line(message):
    sys.stdout.write(f"\r\x1b[2K{terminal_clickable(message)}\n")
    sys.stdout.flush()


def print_line(message):
    sys.stdout.write(f"\r\x1b[2K{terminal_clickable(message)}\n")
    sys.stdout.flush()


def prompt(text, default=None):
    suffix = f" [{default}]" if default else ""
    value = input(f"{text}{suffix}: ").strip()
    if value:
        return value
    return default or ""


def prompt_existing_report_action():
    while True:
        value = input(
            "Type: new / edit / update / update [keyword] / resync [keyword]: "
        ).strip().lower()
        cache = "--cache" in value
        value = value.replace("--cache", "").strip()
        if value in {"new"}:
            return ("new", cache, None)
        if value == "edit":
            return ("edit", False, None)
        if value in {"q", "quit", "exit"}:
            return ("quit", False, None)
        if value in {"update", "upd"}:
            return ("update", cache, None)
        if value.startswith("update ") or value.startswith("upd "):
            pattern = value.split(" ", 1)[1].strip()
            return ("update", cache, pattern)
        if value == "resync":
            return ("resync", False, None)
        if value.startswith("resync "):
            pattern = value.split(" ", 1)[1].strip()
            return ("resync", False, pattern)
        print("Please enter new, edit, update, resync, or quit.")


def prompt_list(text, default=None):
    raw = prompt(text, default)
    if raw.strip().lower() in BACK_COMMANDS:
        return "/back"
    if not raw:
        return []
    return [item.strip() for item in raw.split(",") if item.strip()]


def split_keywords(raw):
    text = (raw or "").strip()
    if not text:
        return []
    if "," in text:
        return [item.strip() for item in text.split(",") if item.strip()]
    return [text]


def sanitize_keyword_values(values):
    if values is None:
        return []
    if isinstance(values, str):
        values = split_keywords(values)
    cleaned = []
    seen = set()
    for value in values:
        item = str(value).strip().lstrip("+-")
        lowered = item.lower()
        if not item:
            continue
        if lowered in BACK_COMMANDS:
            continue
        if lowered.startswith("/"):
            continue
        if lowered.startswith(MALFORMED_EDIT_PREFIXES):
            continue
        if item not in seen:
            cleaned.append(item)
            seen.add(item)
    return cleaned


def parse_edit_operations(text):
    text = (text or "").strip()
    if not text:
        return []
    if text[0] not in "+-":
        return []
    ops = []
    current_sign = None
    for chunk in [part.strip() for part in text.split(",") if part.strip()]:
        if chunk[0] in "+-":
            current_sign = chunk[0]
            value = chunk[1:].strip()
        else:
            value = chunk
        if not value or current_sign is None:
            return []
        ops.append((current_sign, value))
    return ops


def keyword_command_result(raw, current_values):
    text = (raw or "").strip()
    lowered = text.lower()
    if lowered in BACK_COMMANDS:
        return "/back"
    if lowered in CLEAR_COMMANDS:
        return "/clear"
    if not text:
        return []
    if lowered.startswith(MALFORMED_EDIT_PREFIXES):
        print("Use /e +new keyword -old keyword.")
        return None

    head, _, tail = text.partition(" ")
    head = head.lower()
    tail = tail.strip()

    if head in EDIT_COMMANDS:
        text = tail  # strip the /e prefix before parsing

    if text and text[0] in "+-":
        operations = parse_edit_operations(text)
        if not operations:
            print("Use +new keyword, kwd2 to add or -old keyword, kwd2 to remove.")
            return None
        current = list(current_values or [])
        for sign, value in operations:
            if sign == "+":
                for item in split_keywords(value):
                    if item not in current:
                        current.append(item)
            elif sign == "-":
                removals = set(split_keywords(value))
                current = [item for item in current if item not in removals]
        return sanitize_keyword_values(current)

    if head.startswith("/"):
        print("I do not know that command here. Use +kwd to add, -kwd to remove, or /clear to empty the list.")
        return None

    return sanitize_keyword_values(split_keywords(text))


def _edit_hint(noun="kwd"):
    return f"→ +{noun}, {noun} adds; -{noun}, {noun} removes; +{noun}, -{noun} mixes; new list replaces; Enter to skip"


_EDIT_HINT = _edit_hint("kwd")


def prompt_list_or_default(text, default_values=None, required=False, empty_message=None, noun="kwd"):
    if isinstance(default_values, str):
        default_values = [item.strip() for item in default_values.split(",") if item.strip()]
    print(f"{text}:")
    if default_values:
        for kw in default_values:
            print(f"  {kw}")
    else:
        print("  (empty)")
    while True:
        raw = input(f"{_edit_hint(noun)}: ").strip()
        values = keyword_command_result(raw, default_values)
        if values is None:
            continue
        if values == "/back":
            return "/back"
        if values == "/clear":
            if required:
                print(empty_message or "Please enter at least one value.")
                continue
            return []
        if values:
            return values
        if default_values:
            return default_values
        if not required:
            return []
        print(empty_message or "Please enter at least one value.")


def prompt_required_list(text):
    while True:
        raw = prompt(text)
        values = keyword_command_result(raw, None)
        if values is None:
            continue
        if values == "/back":
            return "/back"
        if values == "/clear":
            print("Please enter at least one keyword. I use it to find the epics.")
            continue
        if values:
            return values
        print("Please enter at least one keyword. I use it to find the epics.")


def prompt_optional_list(text):
    while True:
        raw = prompt(text)
        values = keyword_command_result(raw, None)
        if values is None:
            continue
        if values == "/clear":
            return []
        return values


def prompt_drive_folder_url(text):
    while True:
        value = prompt(text)
        if value.strip().lower() in BACK_COMMANDS:
            return "/back"
        if value:
            return value.strip()
        print("Please paste a Google Drive folder URL so I know where to upload the file.")


def prompt_google_oauth_json():
    for line in load_oauth_setup_steps():
        print(line)
    lines = []
    while True:
        line = input()
        if line.strip().lower() in BACK_COMMANDS:
            return "/back"
        if not lines:
            candidate = line.strip()
            if candidate:
                try:
                    data = json.loads(candidate)
                except json.JSONDecodeError:
                    data = None
                if isinstance(data, dict):
                    target = APP_DIR / "google-oauth-client-secrets.json"
                    target.write_text(json.dumps(data, indent=2))
                    return target
        if line.strip() == "END":
            break
        lines.append(line)
    raw = "\n".join(lines).strip()
    if not raw:
        print("Please paste the Google JSON first.")
        return prompt_google_oauth_json()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        print("That did not look like valid JSON. Please paste the full Google JSON again.")
        return prompt_google_oauth_json()
    target = APP_DIR / "google-oauth-client-secrets.json"
    target.write_text(json.dumps(data, indent=2))
    return target


def prompt_yes_no(text, default="y"):
    while True:
        value = prompt(text, default).strip().lower()
        if value in BACK_COMMANDS:
            return "/back"
        if value in {"y", "yes"}:
            return True
        if value in {"n", "no"}:
            return False
        print("Please enter y or n.")


def prompt_yes_no_or_default(text, default_value=None):
    default_text = "y" if default_value is True else "n" if default_value is False else None
    while True:
        value = prompt(text, default_text).strip().lower()
        if value in BACK_COMMANDS:
            return "/back"
        if not value and default_value is not None:
            return default_value
        if value in {"y", "yes"}:
            return True
        if value in {"n", "no"}:
            return False
        print("Please enter y or n.")


def prompt_text_or_default(text, default_value=None, required=False, skip_message=None):
    while True:
        value = prompt(text, default_value if default_value else None)
        if value.strip().lower() in BACK_COMMANDS:
            return "/back"
        if value:
            return value.strip()
        if default_value:
            return str(default_value)
        if not required:
            return ""
        print(skip_message or "Please enter a value.")


def sanitize_expected_tasks_per_week(values):
    cleaned = {}
    for key, value in (values or {}).items():
        feature = str(key).strip()
        if not feature:
            continue
        try:
            number = float(value)
        except Exception:
            continue
        if number > 0:
            cleaned[feature] = number
    return cleaned


def parse_freeform_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    current_year = dt.date.today().year
    formats = [
        "%b %d, %Y",
        "%d %b %Y",
        "%d-%b-%Y",
        "%d-%b-%y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%m-%d-%Y",
        "%b %d %Y",
        "%B %d, %Y",
        "%d %B %Y",
        "%B %d %Y",
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%b %d",
        "%d %b",
        "%B %d",
        "%d %B",
        "%m-%d",
        "%d-%m",
        "%m/%d",
        "%d/%m",
    ]
    for fmt in formats:
        try:
            parsed = dt.datetime.strptime(text, fmt)
        except Exception:
            continue
        if "%Y" not in fmt:
            parsed = parsed.replace(year=current_year)
        return parsed.date()
    return None


def sanitize_feature_eta_dates(values):
    cleaned = {}
    for key, value in (values or {}).items():
        feature = str(key).strip()
        if not feature:
            continue
        parsed = parse_freeform_date(value)
        if parsed:
            cleaned[feature] = parsed.isoformat()
    return cleaned


def _lookup_icase(mapping, key):
    """Case-insensitive dict lookup; returns (found_key, value) or (None, None)."""
    target = str(key or "").strip().lower()
    for k, v in (mapping or {}).items():
        if str(k).strip().lower() == target:
            return k, v
    return None, None


def make_project_relative(path_value):
    if not path_value:
        return None
    path = Path(path_value).expanduser()
    try:
        resolved = path.resolve()
    except Exception:
        resolved = path
    try:
        return str(resolved.relative_to(ROOT))
    except Exception:
        return str(path_value)


def resolve_project_path(path_value):
    if not path_value:
        return None
    path = Path(path_value).expanduser()
    if not path.is_absolute():
        path = (ROOT / path).resolve()
    return path


def pace_or_eta_label(feature):
    return f"Expected tasks per week (number) or ETA date (DD-MMM-YYYY) for {feature}"


def prompt_rate_or_date(text, default_rate=None, default_date=None):
    default_text = None
    if default_rate not in (None, ""):
        default_text = str(int(default_rate)) if float(default_rate).is_integer() else str(default_rate)
    elif default_date:
        default_text = default_date
    suffix = f" [{default_text}]" if default_text else ""
    print(f"{text}{suffix}")
    while True:
        value = input("→ New value / Enter to skip: ").strip()
        if value.lower() in BACK_COMMANDS:
            return "/back"
        if not value:
            if default_rate not in (None, ""):
                return ("rate", float(default_rate))
            if default_date:
                return ("date", default_date)
            return None
        try:
            number = float(value)
            if number <= 0:
                raise ValueError
            return ("rate", number)
        except Exception:
            parsed = parse_freeform_date(value)
            if not parsed:
                print("Please enter a positive number or a date like 11-Jul-2026 or Jul 11, 2026.")
                continue
            return ("date", parsed.isoformat())


def has_full_disk_access():
    """Return True if the current process has Full Disk Access."""
    try:
        result = subprocess.run(
            ["sqlite3",
             os.path.expanduser("~/Library/Application Support/com.apple.TCC/TCC.db"),
             "SELECT 1"],
            capture_output=True,
        )
        return result.returncode == 0
    except Exception:
        return False


def ensure_full_disk_access():
    """Check FDA and guide the user to grant it if missing. Returns True if granted."""
    if has_full_disk_access():
        return True
    print()
    print("⚠️  Full Disk Access is required for automatic background updates.")
    print("   Without it, macOS silently blocks the scheduled job.")
    print()
    print("   Opening System Settings → Privacy & Security → Full Disk Access...")
    print("   Add your Terminal app there, then come back and press Enter.")
    print()
    subprocess.run(["open", "x-apple.systempreferences:com.apple.preference.security?Privacy_AllFiles"])
    input("Press Enter once you've added Terminal to Full Disk Access: ")
    if has_full_disk_access():
        print("✓ Full Disk Access confirmed.")
        return True
    print("⚠️  Still not detected. Automatic updates may not fire until this is granted.")
    return False


def keychain_has_jira():
    cmd = [
        "security",
        "find-generic-password",
        "-a",
        JIRA_ACCOUNT,
        "-s",
        JIRA_SERVICE,
        "-w",
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    return result.returncode == 0 and bool(result.stdout.strip())


def read_jira_token():
    env_token = os.environ.get("JIRA_TOKEN", "").strip()
    if env_token:
        return env_token
    cmd = [
        "security",
        "find-generic-password",
        "-a",
        JIRA_ACCOUNT,
        "-s",
        JIRA_SERVICE,
        "-w",
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode == 0:
        token = result.stdout.strip()
        if token:
            return token
    return ""


def jira_token_is_valid(token):
    if not token:
        return False
    # Read live — the host may have only just been collected during this
    # same run (see run_jira_setup()), so a cached value would be stale.
    host = os.environ.get("JIRA_HOST", "")
    if not host:
        # A malformed "https:///..." request raises urllib.error.URLError,
        # which the generic exception handler below treats as "probably a
        # VPN hiccup, assume the token's fine" — that let an account with
        # an existing cached token but no recorded host skip run_jira_setup()
        # forever, so jira-report.py's own subprocess kept hitting the same
        # empty-host crash. An empty host is never "probably fine".
        return False
    req = urllib.request.Request(
        f"https://{host}/rest/api/2/myself",
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as response:
            return response.status == 200
    except urllib.error.HTTPError as exc:
        if exc.code in {401, 403}:
            return False
        return True  # server reachable but unexpected error — assume token is fine
    except Exception:
        return True  # network/VPN not reachable — assume token is fine, let report fail properly


def is_jira_network_error_output(text):
    lower = (text or "").lower()
    return any(
        token in lower
        for token in [
            "likely vpn timeout",
            "connection appears to have dropped",
            "nodename nor servname provided",
            "temporary failure in name resolution",
            "timed out",
            "connection reset",
            "network is unreachable",
        ]
    )


def daily_job_exists():
    system = sys.platform.lower()
    if system == "darwin":
        result = subprocess.run(["launchctl", "list"], capture_output=True, text=True)
        return JOB_NAME in result.stdout
    if system.startswith("win"):
        result = subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-Command",
                f"if (Get-ScheduledTask -TaskName '{JOB_NAME}' -ErrorAction SilentlyContinue) {{ exit 0 }} else {{ exit 1 }}",
            ],
            cwd=ROOT,
        )
        return result.returncode == 0
    return False


def _trim_to_hostname(value):
    """Accept a bare hostname or a full URL/link (e.g. a pasted task or
    board link) and return just the hostname."""
    value = (value or "").strip()
    if "://" in value:
        value = value.split("://", 1)[1]
    return value.split("/", 1)[0]


_ATLASSIAN_DC_MCP_CONFIG = Path.home() / ".atlassian-dc-mcp" / "jira.env"
_JIRA_TOKEN_HELP = "In Jira: click your avatar → Profile → Personal Access Tokens → Create"


def _write_jira_setup_config(host):
    """(Over)write the Atlassian DC MCP config file with what the user just
    entered. This replaces that tool's own interactive `setup` step (see
    run_jira_setup()) rather than just pre-filling defaults for it, so it
    always writes — unlike a pre-seed, there's no "don't touch an existing
    file" concern here, since this *is* the setup now."""
    try:
        _ATLASSIAN_DC_MCP_CONFIG.parent.mkdir(parents=True, exist_ok=True)
        _ATLASSIAN_DC_MCP_CONFIG.write_text(
            f"JIRA_HOST={host}\nJIRA_API_BASE_PATH=/rest\nJIRA_DEFAULT_PAGE_SIZE=25\n"
        )
    except Exception:
        pass  # best-effort — the config file is a convenience, not required for this tool's own use


def _store_jira_token_in_keychain(token):
    subprocess.run(
        ["security", "add-generic-password", "-a", JIRA_ACCOUNT, "-s", JIRA_SERVICE, "-w", token, "-U"],
        capture_output=True, text=True,
    )


def _check_jira_reachability(host, token):
    """Distinguish "not on VPN" from "bad token" from "working" for the
    setup flow itself. Unlike jira_token_is_valid() (used elsewhere, where
    a network error optimistically means "assume the token is fine, let a
    later step surface VPN issues its own way"), here we want to say
    plainly what's wrong right now, since the user is watching and waiting."""
    req = urllib.request.Request(
        f"https://{host}/rest/api/2/myself",
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as response:
            return "ok" if response.status == 200 else "error"
    except urllib.error.HTTPError as exc:
        return "invalid_token" if exc.code in (401, 403) else "error"
    except urllib.error.URLError:
        return "unreachable"
    except Exception:
        return "error"


def _openpyxl_importable():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def _ensure_openpyxl():
    """jira-report.py needs openpyxl to build the Excel file, but it fails
    at its own top-level import if missing — same chicken-and-egg problem
    as Python itself not being able to check for Python — so check here,
    under the same python executable jira-report.py will actually run
    under, before ever invoking it."""
    if _openpyxl_importable():
        return True
    print("openpyxl is not installed, and we need it to build the report.")
    answer = input("Do you agree to install it? (y/n): ").strip().lower()
    if answer not in ("y", "yes"):
        return False
    print("Installing openpyxl...")
    # Deliberately not captured/spinner-wrapped — pip prints its own live
    # download progress, and capturing it would just hide that percentage
    # from the user until the whole thing finished (same reason interactive
    # subprocesses never get wrapped in our own spinner elsewhere).
    result = subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"])
    if result.returncode != 0:
        print("Couldn't install openpyxl — see the error above.")
        return False
    return True


def _persist_jira_host(host):
    """Save the collected host to settings so future runs (fresh Python
    processes, no memory of this one) can restore it into os.environ
    automatically — see main()'s startup — instead of relying on the user
    manually adding it to their shell profile, which is exactly the step
    that silently didn't happen and caused this to regress once already."""
    try:
        current = load_state()
        current["jira_host"] = host
        save_state(current)
    except Exception:
        pass  # best-effort — os.environ is already set for this run either way


def run_jira_setup():
    print("\nJira is not set up on this computer yet.")
    host = ""
    while not host:
        host = _trim_to_hostname(input("→ Enter Jira URL: "))
    token = ""
    while not token:
        token = input(f"→ Enter Jira Personal Access Token ({_JIRA_TOKEN_HELP}): ").strip()

    def _finish_setup():
        os.environ["JIRA_HOST"] = host
        _persist_jira_host(host)
        _write_jira_setup_config(host)
        _store_jira_token_in_keychain(token)
        return _check_jira_reachability(host, token)

    status = run_spinner("Setting things up — Jira and your token are getting acquainted...", _finish_setup)

    if status == "unreachable":
        raise SystemExit(
            f"\nCan't reach Jira — connect to VPN and check that https://{host} opens in your "
            "browser, then run this tool again."
        )
    if status == "invalid_token":
        raise SystemExit("\nThat token doesn't look valid. Double-check it and run this tool again.")
    if status != "ok":
        raise SystemExit("\nCouldn't confirm Jira is reachable. Run this tool again to retry.")

    print("✓ Jira is set up. This is saved for future runs — nothing else to do.")


def _check_existing_jira_token():
    token = read_jira_token()
    return token if token and jira_token_is_valid(token) else None


def ensure_jira_token():
    # Only the fast, silent check runs under a spinner. run_jira_setup() itself
    # prompts the user directly (host, then token) via input() — a spinner
    # redrawing over stdout every 120ms would fight with those prompts the
    # same way it once did with the third-party npx wizard this replaced.
    token = run_spinner("Checking Jira setup. Making sure the gears are greased...", _check_existing_jira_token)
    if token:
        return token
    run_jira_setup()
    token = run_spinner("Verifying Jira setup...", _check_existing_jira_token)
    if token:
        return token
    raise SystemExit(
        "Jira setup did not produce a reusable token. Please re-run setup and make sure the token is saved."
    )


def load_state():
    SETTINGS_DIR.mkdir(exist_ok=True)
    if not STATE_PATH.exists() and SETTINGS_TEMPLATE.exists():
        STATE_PATH.write_text(SETTINGS_TEMPLATE.read_text())
    if STATE_PATH.exists():
        state = json.loads(STATE_PATH.read_text())
        # Normalise new `features` list → legacy flat keys
        if "features" in state and not state.get("include"):
            fl = state["features"]
            state["include"] = [f["keyword"] for f in fl if f.get("keyword")]
            if "expected_tasks_per_week" not in state:
                state["expected_tasks_per_week"] = {
                    f["keyword"]: f["expected_pace"]
                    for f in fl if f.get("keyword") and f.get("expected_pace") is not None
                }
            if "feature_eta_dates" not in state:
                state["feature_eta_dates"] = {
                    f["keyword"]: f["eta"]
                    for f in fl if f.get("keyword") and f.get("eta")
                }
        state["include"] = _strip_placeholders(sanitize_keyword_values(state.get("include")))
        state["exclude"] = _strip_placeholders(sanitize_keyword_values(state.get("exclude")))
        state["expected_tasks_per_week"] = sanitize_expected_tasks_per_week(state.get("expected_tasks_per_week"))
        state["feature_eta_dates"] = sanitize_feature_eta_dates(state.get("feature_eta_dates"))
        output = state.get("output")
        if output:
            state["output"] = str(Path(output).name if Path(output).is_absolute() else Path(output))
        if state.get("google_client_secrets"):
            state["google_client_secrets"] = str(resolve_project_path(state["google_client_secrets"]))
        return state
    return {}


_PLACEHOLDER_KWDS = {"kwd", "kwd1", "kwd2", "kwd3"}


def _strip_placeholders(values):
    return [v for v in (values or []) if v.lower() not in _PLACEHOLDER_KWDS]


def save_state(state):
    state = dict(state)
    if "exclude" in state:
        state["exclude"] = _strip_placeholders(state["exclude"])
    if "include" in state:
        state["include"] = _strip_placeholders(state["include"])
    if "features" in state:
        state["features"] = [f for f in state["features"] if f.get("keyword", "").lower() not in _PLACEHOLDER_KWDS]
    SETTINGS_DIR.mkdir(exist_ok=True)
    STATE_PATH.write_text(json.dumps(state, indent=2))


def parse_update_time(raw):
    """Parse 'HH:MM' from user input. Returns (h, m) or None if invalid."""
    raw = (raw or "").strip()
    try:
        parts = raw.split(":")
        h, m = int(parts[0]), int(parts[1])
        if 0 <= h <= 23 and 0 <= m <= 59:
            return h, m
    except Exception:
        pass
    return None


def _detect_system_timezone():
    """Best-effort local timezone, no extra dependency required: the IANA
    name if /etc/localtime resolves to one (handles DST correctly), else
    the current UTC offset (parse_timezone in run-daily-update.py accepts
    either form)."""
    try:
        resolved = os.path.realpath("/etc/localtime")
        if "zoneinfo/" in resolved:
            return resolved.split("zoneinfo/", 1)[1]
    except Exception:
        pass
    offset = dt.datetime.now().astimezone().utcoffset()
    if offset is None:
        return "UTC"
    total_minutes = int(offset.total_seconds() // 60)
    sign = "+" if total_minutes >= 0 else "-"
    hours, minutes = divmod(abs(total_minutes), 60)
    return f"{sign}{hours:02d}:{minutes:02d}"


def save_draft_state(include_keywords=None, excludes=None, project_keys=None, output=None,
                     local_only=None, drive_folder=None, google_client_secrets=None,
                     expected_tasks_per_week=None, feature_eta_dates=None,
                     update_time=None, update_timezone=None, auto_update=None,
                     done_statuses=None):
    current = load_state()
    if include_keywords is not None:
        current["include"] = sanitize_keyword_values(include_keywords)
    else:
        current["include"] = _strip_placeholders(current.get("include", []))
    if excludes is not None:
        current["exclude"] = sanitize_keyword_values(excludes)
    else:
        current["exclude"] = _strip_placeholders(current.get("exclude", []))
    if project_keys is not None:
        current["project_keys"] = [k.strip().upper() for k in project_keys if k.strip()]
    if done_statuses is not None:
        current["done_statuses"] = [s.strip() for s in done_statuses if s.strip()]
    if expected_tasks_per_week is not None:
        current["expected_tasks_per_week"] = sanitize_expected_tasks_per_week(expected_tasks_per_week)
    if feature_eta_dates is not None:
        current["feature_eta_dates"] = sanitize_feature_eta_dates(feature_eta_dates)
    if output is not None:
        current["output"] = str(output)
    if local_only is not None:
        current["local_only"] = local_only
    if drive_folder is not None:
        current["drive_folder"] = str(drive_folder)
    if google_client_secrets is not None:
        current["google_client_secrets"] = make_project_relative(google_client_secrets)
    if update_time is not None:
        current["update_time"] = update_time
    if update_timezone is not None:
        current["update_timezone"] = update_timezone
    if auto_update is not None:
        current["auto_update"] = auto_update
    save_state(current)


def _run_streaming_and_capture(cmd, cwd, env, startup_message=None):
    """Run cmd, relaying its combined stdout+stderr live while also capturing
    everything, so it can be saved to the crash log if the process fails with
    something we don't already handle explicitly.

    Reads with read1() rather than iterating by line: jira-report.py's own
    spinner redraws a line with a bare \\r (no trailing \\n) between frames,
    and `for line in proc.stdout` blocks until it sees a \\n — so the spinner
    frames would sit invisible in the pipe until the next newline-terminated
    write (a "✓ ..." line) flushed them all at once, reading as a freeze.

    If startup_message is given, a spinner shows it until the subprocess's
    first byte of output arrives, covering the silent gap while the child
    process starts up and connects before it prints anything of its own.
    """
    buf = []
    proc = subprocess.Popen(cmd, cwd=cwd, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)

    stop = threading.Event()
    first_chunk = threading.Event()
    thread = None
    if startup_message:
        def animate():
            idx = 0
            while not stop.is_set() and not first_chunk.is_set():
                sys.stdout.write(spinner_line(SYMBOLS[idx % len(SYMBOLS)], startup_message))
                sys.stdout.flush()
                idx += 1
                stop.wait(0.12)

        thread = threading.Thread(target=animate, daemon=True)
        thread.start()

    while True:
        chunk = proc.stdout.read1(4096)
        if not chunk:
            break
        if not first_chunk.is_set():
            first_chunk.set()
            if thread:
                thread.join(timeout=1)
        sys.stdout.buffer.write(chunk)
        sys.stdout.buffer.flush()
        buf.append(chunk.decode(errors="replace"))
    stop.set()
    if thread:
        thread.join(timeout=1)
    proc.wait()
    return proc.returncode, "".join(buf)


def run_spinner(message, work_fn):
    stop = threading.Event()

    def animate():
        idx = 0
        while not stop.is_set():
            sys.stdout.write(spinner_line(SYMBOLS[idx % len(SYMBOLS)], message))
            sys.stdout.flush()
            idx += 1
            stop.wait(0.12)

    thread = threading.Thread(target=animate, daemon=True)
    thread.start()
    try:
        return work_fn()
    finally:
        stop.set()
        thread.join(timeout=1)
        sys.stdout.write(f"\r\x1b[2K✓ {message}\n")
        sys.stdout.flush()


def load_oauth_setup_steps():
    if OAUTH_SETUP_FILE.exists():
        return json.loads(OAUTH_SETUP_FILE.read_text()).get("steps", [])
    return []


def slugify(value):
    slug = "".join(ch.lower() if ch.isalnum() else "-" for ch in value)
    slug = "-".join(part for part in slug.split("-") if part)
    return slug or "roadmap"


def current_output_path():
    return Path("report") / f"roadmap {dt.date.today().year}.xlsx"


def find_google_client_secrets():
    env_path = os.environ.get("GOOGLE_OAUTH_CLIENT_SECRETS", "").strip()
    if env_path:
        candidate = Path(env_path).expanduser()
        if candidate.exists():
            return candidate
    for candidate in GOOGLE_CLIENT_SECRETS_CANDIDATES:
        if candidate.exists() and not is_placeholder_client_config(candidate):
            return candidate
    return None


def _do_drive_upload(output, local_only, drive_folder, google_client_secrets):
    if local_only or not drive_folder or not google_client_secrets:
        return
    run_spinner(
        "Redirecting to drive.google.com. Awaiting sign-in access...",
        lambda: authorize(google_client_secrets),
    )
    status_line("✓ Access granted.")
    upload_result = {}
    def _upload_all():
        upload_result["value"] = upload_or_update(output, drive_folder, Path(output).name, google_client_secrets)
        if STATE_PATH.exists():
            upload_or_update(STATE_PATH, drive_folder, "roadmap-settings.json", google_client_secrets)
    run_spinner("Uploading roadmap and local settings to Google Drive...", _upload_all)
    result = upload_result.get("value", {}) or {}
    link = result.get("webViewLink") or drive_folder
    status_line(f"✓ Roadmap and settings uploaded to Google Drive folder: {drive_folder}")
    status_line("Open roadmap:")
    print_line(link)


def _prompt_edit_section():
    # "resync a feature" isn't offered here — it's its own top-level command
    # (`resync` / `resync [keyword]` at the main prompt), not a settings
    # field to edit, so it doesn't belong in this list.
    _OPTIONS = {"all", "keyword", "keywords", "jira key", "jira keys", "jira", "keyword eta", "eta", "url", "drive", "time", "timezone", "status", "statuses", "done status", "done statuses", "done"}
    while True:
        value = input("Edit: all / keyword / Jira key / keyword eta / URL / time / Done status: ").strip().lower()
        if not value or value == "all":
            return "all"
        if value in {"keyword", "keywords"}:
            return "keyword"
        if value in {"jira key", "jira keys", "jira"}:
            return "jira key"
        if value in {"keyword eta", "eta"}:
            return "keyword eta"
        if value in {"url", "drive"}:
            return "url"
        if value in {"time", "timezone"}:
            return "time"
        if value in {"status", "statuses", "done status", "done statuses", "done"}:
            return "status"
        print("Enter: all, keyword, Jira key, keyword eta, URL, time, or Done status.")


def _run_targeted_edit(
    section,
    include_keywords, excludes, project_keys,
    expected_tasks_per_week, feature_eta_dates,
    local_only, drive_folder, google_client_secrets,
    update_time, update_timezone, auto_update=True,
    done_statuses=None,
):
    done_statuses = done_statuses or DEFAULT_DONE_STATUSES
    if section == "keyword":
        new_include = prompt_list_or_default(
            "Keywords to include, comma-separated",
            include_keywords or None,
            required=True,
            empty_message="Please enter at least one keyword.",
        )
        if new_include not in ("/back", None):
            include_keywords = sanitize_keyword_values(new_include)
        new_exclude = prompt_list_or_default(
            "Keywords to exclude, comma-separated (Enter to skip)",
            excludes or None,
        )
        if new_exclude not in ("/back", None):
            excludes = sanitize_keyword_values(new_exclude)
        new_keys = prompt_list_or_default(
            "Jira project keys (the prefix before the dash in a Jira issue ID, e.g. ABC in ABC-123) — empty for all",
            project_keys or None,
        )
        if new_keys not in ("/back", None):
            project_keys = [k.strip().upper() for k in new_keys if k.strip()]

    elif section == "jira key":
        new_keys = prompt_list_or_default(
            "Jira project keys (the prefix before the dash in a Jira issue ID, e.g. ABC in ABC-123) — empty for all",
            project_keys or None,
        )
        if new_keys not in ("/back", None):
            project_keys = [k.strip().upper() for k in new_keys if k.strip()]

    elif section == "keyword eta":
        features_with_values = [
            kw for kw in include_keywords
            if kw in expected_tasks_per_week or kw in feature_eta_dates
        ]
        if not features_with_values:
            features_with_values = include_keywords
        for feature in features_with_values:
            result = prompt_rate_or_date(
                pace_or_eta_label(feature),
                expected_tasks_per_week.get(feature),
                feature_eta_dates.get(feature),
            )
            if result == "/back":
                break
            if result is None:
                continue
            kind, stored = result
            if kind == "rate":
                expected_tasks_per_week[feature] = stored
                feature_eta_dates.pop(feature, None)
            else:
                feature_eta_dates[feature] = stored
                expected_tasks_per_week.pop(feature, None)

    elif section == "url":
        default_sync = (not local_only) if local_only is not None else True
        default_text = "y" if default_sync else "n"
        print(f"Sync with Google Drive (y/n) [{default_text}]:")
        raw_sync = input("→ New value / Enter to skip: ").strip().lower()
        if raw_sync in {"y", "yes"}:
            local_only = False
        elif raw_sync in {"n", "no"}:
            local_only = True
        if not local_only:
            print(f"Google Drive folder URL{f' [{drive_folder}]' if drive_folder else ''}:")
            raw_url = input("→ New URL / Enter to skip: ").strip()
            if raw_url:
                drive_folder = raw_url
            if not google_client_secrets:
                google_client_secrets = find_google_client_secrets()

    elif section == "time":
        default_au = "y" if auto_update else "n"
        print(f"Automatic updates? (y/n) [{default_au}]:")
        raw_au = input("→ New value / Enter to skip: ").strip().lower()
        if raw_au in {"y", "yes"}:
            auto_update = True
        elif raw_au in {"n", "no"}:
            auto_update = False
        if not auto_update:
            if sys.platform == "darwin":
                subprocess.run(["bash", str(APP_DIR / "uninstall-launchd.sh")], cwd=ROOT, check=False)
                print("✓ Automatic updates removed.")
        else:
            print(f"Automatic daily update time (HH:MM in 24h){f' [{update_time}]' if update_time else ''}:")
            raw_time = input("→ New time / Enter to skip: ").strip()
            if raw_time and raw_time not in {"/back", "/b"}:
                parsed = parse_update_time(raw_time)
                if parsed:
                    update_time = f"{parsed[0]:02d}:{parsed[1]:02d}"
                else:
                    print("Please enter a time in HH:MM format, e.g. 09:00.")
            print(f"Timezone (e.g. Europe/Kyiv, America/New_York, UTC){f' [{update_timezone}]' if update_timezone else ''}:")
            raw_tz = input("→ New timezone / Enter to skip: ").strip()
            if raw_tz and raw_tz not in {"/back", "/b"}:
                update_timezone = raw_tz

    elif section == "status":
        print("Which Jira statuses count as done — no more active dev work you would like to track towards ETA?")
        new_statuses = prompt_list_or_default(
            "Currently Done is",
            done_statuses or DEFAULT_DONE_STATUSES,
            required=True,
            empty_message="Please enter at least one status.",
            noun="status",
        )
        if new_statuses not in ("/back", None):
            done_statuses = [s.strip() for s in new_statuses if s.strip()]

    # Save merged state
    _state = load_state()
    eta_map = dict(sanitize_feature_eta_dates(feature_eta_dates) or {})
    pace_map = dict(sanitize_expected_tasks_per_week(expected_tasks_per_week) or {})
    features_out = []
    for kw in sanitize_keyword_values(include_keywords):
        entry = {"keyword": kw}
        _, eta_val = _lookup_icase(eta_map, kw)
        if eta_val is not None:
            entry["eta"] = eta_val
        _, pace_val = _lookup_icase(pace_map, kw)
        if pace_val is not None:
            entry["expected_pace"] = pace_val
        features_out.append(entry)
    _state["features"] = features_out
    _state["exclude"] = sanitize_keyword_values(excludes)
    _state["project_keys"] = [k.strip().upper() for k in project_keys if k.strip()]
    _state["done_statuses"] = [s.strip() for s in (done_statuses or DEFAULT_DONE_STATUSES) if s.strip()]
    _state["local_only"] = local_only
    if drive_folder:
        _state["drive_folder"] = drive_folder
    if google_client_secrets:
        _state["google_client_secrets"] = make_project_relative(google_client_secrets)
    if update_time:
        _state["update_time"] = update_time
    if update_timezone:
        _state["update_timezone"] = update_timezone
    _state["auto_update"] = auto_update
    save_state(_state)


def _resync_feature_flow(include_keywords, keyword=None):
    """Move one existing keyword's epics' tasks into their own feature
    without touching anything else in the report — e.g. an epic that used
    to only match a generic catch-all keyword got renamed/re-scoped to
    match a more specific one, and its tasks are stuck under the old label.
    `--update --new-features <keyword>` already does exactly this (see
    jira-report.py), but was only ever reachable by invoking jira-report.py
    directly — this exposes it from the interactive menu (`resync` /
    `resync [keyword]` at the main prompt) instead.

    If `keyword` isn't given (bare "resync"), prompts with a picker over
    configured keywords. If it IS given ("resync [keyword]"), it's passed
    straight through without validating it against configured keywords —
    same as "update [keyword]" already does — so jira-report.py's own epic
    matching is the source of truth, not a second copy of it here."""
    if keyword:
        resync_keyword = keyword
    else:
        if not include_keywords:
            raise SystemExit("No keywords configured yet — nothing to resync.")
        print(
            "Which feature keyword should be resynced? This moves any of its "
            "epics' tasks in from wherever else they're currently sitting, "
            "without touching anything else."
        )
        for kw in include_keywords:
            print(f"  {kw}")
        while True:
            chosen = input("→ Keyword to resync (Enter to cancel): ").strip()
            if not chosen or chosen.lower() in BACK_COMMANDS:
                raise SystemExit(0)
            matches = [kw for kw in include_keywords if kw.strip().lower() == chosen.lower()]
            if matches:
                break
            print("Enter one of the keywords listed above exactly.")
        resync_keyword = matches[0]

    ensure_jira_token()
    if not _ensure_openpyxl():
        raise SystemExit(
            f"openpyxl is required to build the report. Install it yourself with:\n"
            f"    {sys.executable} -m pip install openpyxl\nthen run this tool again."
        )

    output = str(current_output_path())
    cmd = [
        sys.executable,
        str(APP_DIR / "jira-report.py"),
        "--state", str(STATE_PATH),
        "--output", output,
        "--update",
        "--new-features", resync_keyword,
    ]
    env = os.environ.copy()
    jira_token = read_jira_token()
    if jira_token:
        env["JIRA_TOKEN"] = jira_token

    startup_message = f"Resyncing '{resync_keyword}'..."
    try:
        returncode, _proc_out = _run_streaming_and_capture(cmd, ROOT, env, startup_message=startup_message)
    except KeyboardInterrupt:
        raise SystemExit(0)
    if returncode == 86:
        print_line("Jira sign-in looks stale. Let’s refresh setup once.")
        run_jira_setup()
        jira_token = read_jira_token()
        if jira_token:
            env["JIRA_TOKEN"] = jira_token
        try:
            returncode, _proc_out = _run_streaming_and_capture(cmd, ROOT, env, startup_message=startup_message)
        except KeyboardInterrupt:
            raise SystemExit(0)
    if returncode == 87:
        print_line(
            f"Can't reach Jira — connect to VPN and check that https://{env.get('JIRA_HOST', '')} "
            "opens in your browser, then try again."
        )
        raise SystemExit(returncode)
    if returncode not in (0, 88):  # 88 = jira-report.py's "ran fine, nothing changed"
        if "Traceback (most recent call last):" in _proc_out:
            log_path = _write_crash_log(_proc_out)
            print_line(f"Resync crashed. A log was saved to:\n    {log_path}\nPlease send that file so this can be debugged.")
        else:
            print_line("Resync stopped. Please check the keyword and try again.")
        raise SystemExit(returncode)
    print_line(f"✓ '{resync_keyword}' resynced.")


def _send_notification(message: str) -> None:
    """Send a macOS notification via terminal-notifier (no-op on other platforms).

    Belt-and-suspenders PYTEST_CURRENT_TEST check: this has no per-call
    test-mode flag, so any test exercising main()'s update_run path that
    forgets to monkeypatch this out would fire a REAL, live notification —
    that happened here (spoofed as sent by Script Editor, since that's the
    sender id below), and a test that quietly spams the real machine still
    passes, so there was no signal anything was wrong.
    """
    if os.environ.get("PYTEST_CURRENT_TEST"):
        return
    if sys.platform != "darwin":
        return
    import shutil as _shutil
    notifier = _shutil.which("terminal-notifier") or "/opt/homebrew/bin/terminal-notifier"
    if not Path(notifier).exists():
        return
    subprocess.run(
        [notifier, "-message", message, "-title", "Jira Roadmap", "-sender", "com.apple.ScriptEditor2"],
        capture_output=True,
    )


def _restore_jira_host_from_settings(state):
    """A prior run's os.environ change doesn't carry over to this fresh
    process, and relying on the user having added JIRA_HOST to their shell
    profile themselves is exactly what silently didn't happen once already
    — restore it from settings instead if it's not already set."""
    if not os.environ.get("JIRA_HOST") and state.get("jira_host"):
        os.environ["JIRA_HOST"] = state["jira_host"]


def _restore_jira_host_from_external_config():
    """settings/roadmap-settings.local.json lives inside the project folder,
    so deleting/re-cloning the project wipes it — but the Jira token in
    Keychain survives that untouched, since Keychain isn't tied to any
    folder. Without this fallback, a still-valid token gets treated as "not
    set up" purely because JIRA_HOST has nowhere left to come from, and the
    user is asked to redo setup for no reason. ~/.atlassian-dc-mcp/jira.env
    lives outside the project folder (see _write_jira_setup_config), so it
    survives a project folder deletion the same way Keychain does."""
    if os.environ.get("JIRA_HOST"):
        return
    try:
        for line in _ATLASSIAN_DC_MCP_CONFIG.read_text().splitlines():
            if line.startswith("JIRA_HOST="):
                host = line.split("=", 1)[1].strip()
                if host:
                    os.environ["JIRA_HOST"] = host
                return
    except Exception:
        pass  # best-effort — a missing/unreadable file just means no fallback available


def main():
    debug = "--debug" in sys.argv[1:]
    from_cache = "--cache" in sys.argv[1:]
    state = load_state()
    _restore_jira_host_from_settings(state)
    _restore_jira_host_from_external_config()
    REPORTS_DIR.mkdir(exist_ok=True)

    stage = 0
    # Derive include_keywords from features list if present (new format)
    _features = state.get("features") or []
    if _features and not state.get("include"):
        include_keywords = [f["keyword"] for f in _features if f.get("keyword")]
        expected_tasks_per_week = sanitize_expected_tasks_per_week(
            {f["keyword"]: f["expected_pace"] for f in _features if f.get("keyword") and f.get("expected_pace") is not None}
        )
        feature_eta_dates = sanitize_feature_eta_dates(
            {f["keyword"]: f["eta"] for f in _features if f.get("keyword") and f.get("eta")}
        )
    else:
        include_keywords = state.get("include", [])
        expected_tasks_per_week = sanitize_expected_tasks_per_week(state.get("expected_tasks_per_week"))
        feature_eta_dates = sanitize_feature_eta_dates(state.get("feature_eta_dates"))
    excludes = state.get("exclude", [])
    output = str(current_output_path())
    local_only = state.get("local_only")
    drive_folder = state.get("drive_folder")
    google_client_secrets = state.get("google_client_secrets")
    project_keys = state.get("project_keys", [])
    done_statuses = state.get("done_statuses") or DEFAULT_DONE_STATUSES
    update_time = state.get("update_time", "08:00")
    update_timezone = state.get("update_timezone", "UTC")
    auto_update = state.get("auto_update", True)
    fresh_run = False
    update_run = False
    edit_run = False
    filter_pattern = None
    original_keywords_norm = {normalize_keyword(k) for k in include_keywords}
    current_report_file = ROOT / state.get("output", str(current_output_path()))
    _template_content = SETTINGS_TEMPLATE.read_text().strip() if SETTINGS_TEMPLATE.exists() else None
    _local_content = STATE_PATH.read_text().strip() if STATE_PATH.exists() else None
    has_settings = bool(
        _local_content
        and _local_content != _template_content
    )
    if has_settings:
        action, action_cache, filter_pattern = prompt_existing_report_action()
        from_cache = from_cache or action_cache
        if action == "quit":
            raise SystemExit(0)
        if action == "update":
            update_run = True
            stage = 999
        elif action == "resync":
            _resync_feature_flow(include_keywords, filter_pattern)
            raise SystemExit(0)
        elif action == "new":
            fresh_run = True
            print_line("Rebuilding new report using local settings...")
            stage = 999
        elif action == "edit":
            # Route edits through --update (task-level incremental refresh)
            # rather than falling through with neither --fresh nor --update
            # set. That middle-ground path queries Jira for tasks with
            # status NOT IN (done, rejected) — correct for --update, where
            # done tasks are already trusted to be in the existing file —
            # but with nothing else marking the feature as already fully
            # done, that exclusion silently drops real completed work with
            # no fallback to preserve it, e.g. any epic whose tasks are ALL
            # currently Done shows 0 child tasks and vanishes from the report.
            edit_run = True
            edit_section = _prompt_edit_section()
            if edit_section != "all":
                _run_targeted_edit(
                    edit_section,
                    include_keywords, excludes, project_keys,
                    expected_tasks_per_week, feature_eta_dates,
                    local_only, drive_folder, google_client_secrets,
                    update_time, update_timezone, auto_update,
                    done_statuses,
                )
                if edit_section == "time":
                    # time edit only touches the schedule — no report rebuild needed
                    _s = load_state()
                    _new_auto = _s.get("auto_update", True)
                    _new_time = _s.get("update_time", "08:00")
                    _new_tz = _s.get("update_timezone", "UTC")
                    if _new_auto:
                        if sys.platform == "darwin":
                            ensure_full_disk_access()
                            run_spinner(
                                "Updating schedule...",
                                lambda: subprocess.run(["bash", str(APP_DIR / "install-launchd.sh")], cwd=ROOT, check=True),
                            )
                        print(f"✓ Your report will arrive at {_new_time} {_new_tz}, fresh and slightly caffeinated.")
                    raise SystemExit(0)
                _s = load_state()
                include_keywords = [f["keyword"] for f in _s.get("features", []) if f.get("keyword")] or _s.get("include", [])
                excludes = _s.get("exclude", [])
                project_keys = _s.get("project_keys", [])
                done_statuses = _s.get("done_statuses") or DEFAULT_DONE_STATUSES
                local_only = _s.get("local_only", False)
                drive_folder = _s.get("drive_folder", "")
                google_client_secrets_raw = _s.get("google_client_secrets")
                google_client_secrets = str(resolve_project_path(google_client_secrets_raw)) if google_client_secrets_raw else None
                update_time = _s.get("update_time", "08:00")
                update_timezone = _s.get("update_timezone", "UTC")
                auto_update = _s.get("auto_update", True)
                stage = 999
            # edit_section == "all": fall through to wizard (stage 0) with existing values as defaults
    else:
        fresh_run = True
        include_keywords = []
        excludes = []
        project_keys = []
        expected_tasks_per_week = {}
        feature_eta_dates = {}
        local_only = False
        drive_folder = ""
        google_client_secrets = None
        update_time = "08:00"
        update_timezone = _detect_system_timezone()
        auto_update = True

    while True:
        if stage == 999:
            break
        if stage == 0:
            if include_keywords:
                next_include_keywords = prompt_list_or_default(
                    "Keywords to include, comma-separated",
                    include_keywords,
                    required=True,
                    empty_message="Please enter at least one keyword. I use it to find the epics.",
                )
            else:
                next_include_keywords = prompt_required_list("Keywords to include, comma-separated")
            if next_include_keywords == "/back":
                continue
            include_keywords = sanitize_keyword_values(next_include_keywords)
            save_draft_state(include_keywords=include_keywords)
            stage = 1
            continue
        if stage == 1:
            if excludes:
                next_excludes = prompt_list_or_default(
                    "Keywords to exclude, comma-separated (Enter to skip)",
                    excludes,
                    required=False,
                )
            else:
                next_excludes = prompt_optional_list("Keywords to exclude, comma-separated (Enter to skip)")
            if next_excludes == "/back":
                stage = 0
                continue
            excludes = sanitize_keyword_values(next_excludes)
            save_draft_state(excludes=excludes)
            stage = 11
            continue
        if stage == 11:
            if project_keys:
                next_keys = prompt_list_or_default(
                    "Jira project keys to include, comma-separated (the prefix before the dash in a Jira issue ID, e.g. ABC in ABC-123) — leave empty to include all",
                    project_keys,
                    required=False,
                )
            else:
                next_keys = prompt_optional_list("Jira project keys to include, comma-separated (the prefix before the dash in a Jira issue ID, e.g. ABC in ABC-123) — leave empty to include all")
            if next_keys == "/back":
                stage = 1
                continue
            project_keys = [k.strip().upper() for k in next_keys if k.strip()]
            save_draft_state(project_keys=project_keys)
            stage = 12
            continue
        if stage == 12:
            print("Which Jira statuses count as done — no more active dev work you would like to track towards ETA?")
            next_statuses = prompt_list_or_default(
                "Currently Done is",
                done_statuses or DEFAULT_DONE_STATUSES,
                required=True,
                empty_message="Please enter at least one status.",
                noun="status",
            )
            if next_statuses == "/back":
                stage = 11
                continue
            done_statuses = [s.strip() for s in next_statuses if s.strip()]
            save_draft_state(done_statuses=done_statuses)
            stage = 2
            continue
        if stage == 2:
            feature_idx = 0
            while feature_idx < len(include_keywords):
                feature = include_keywords[feature_idx]
                value = prompt_rate_or_date(
                    pace_or_eta_label(feature),
                    expected_tasks_per_week.get(feature),
                    feature_eta_dates.get(feature),
                )
                if value == "/back":
                    if feature_idx == 0:
                        stage = 12
                        break
                    feature_idx -= 1
                    continue
                if value is None:
                    feature_idx += 1
                    continue
                kind, stored = value
                if kind == "rate":
                    expected_tasks_per_week[feature] = stored
                    feature_eta_dates.pop(feature, None)
                else:
                    feature_eta_dates[feature] = stored
                    expected_tasks_per_week.pop(feature, None)
                save_draft_state(expected_tasks_per_week=expected_tasks_per_week, feature_eta_dates=feature_eta_dates)
                feature_idx += 1
            if stage == 1:
                continue
            stage = 3
            continue
        if stage == 3:
            default_sync = (not local_only) if local_only is not None else None
            default_text = "y" if default_sync is True else "n" if default_sync is False else None
            print(f"Sync with Google Drive (y/n){f' [{default_text}]' if default_text else ''}:")
            _raw_sync = input("→ New value / Enter to skip: ").strip().lower()
            if _raw_sync in {"/back", "/b"}:
                stage = 2
                continue
            if _raw_sync in {"y", "yes"}:
                next_sync_with_google_drive = True
            elif _raw_sync in {"n", "no"}:
                next_sync_with_google_drive = False
            elif not _raw_sync and default_sync is not None:
                next_sync_with_google_drive = default_sync
            else:
                print("Please enter y or n.")
                continue
            sync_with_google_drive = next_sync_with_google_drive
            local_only = not sync_with_google_drive
            if sync_with_google_drive:
                stage = 4
                continue
            save_draft_state(local_only=local_only)
            stage = 5
            continue
        if stage == 4:
            print(f"Google Drive folder URL{f' [{drive_folder}]' if drive_folder else ''}:")
            _raw_url = input("→ New URL / Enter to skip: ").strip()
            if _raw_url in {"/back", "/b"}:
                stage = 3
                continue
            if _raw_url:
                drive_folder = _raw_url
            elif not drive_folder:
                print("Please paste a Google Drive folder URL so I know where to upload the file.")
                continue
            google_client_secrets = find_google_client_secrets()
            if not google_client_secrets:
                google_client_secrets = prompt_google_oauth_json()
                if google_client_secrets == "/back":
                    stage = 2
                    continue
            save_draft_state(local_only=local_only, drive_folder=drive_folder, google_client_secrets=google_client_secrets)
            stage = 5
            continue
        if stage == 5:
            default_au = "y" if auto_update else "n"
            print(f"Automatic updates? (y/n) [{default_au}]:")
            raw_au = input("→ New value / Enter to skip: ").strip().lower()
            if raw_au in {"/back", "/b"}:
                stage = 3
                continue
            if raw_au in {"y", "yes"}:
                auto_update = True
            elif raw_au in {"n", "no"}:
                auto_update = False
            elif not raw_au:
                pass  # keep default
            else:
                print("Please enter y or n.")
                continue
            if not auto_update:
                save_draft_state(update_time=update_time, update_timezone=update_timezone, auto_update=auto_update)
                break

            print(f"Automatic daily update time (HH:MM in 24h){f' [{update_time}]' if update_time else ''}:")
            raw_time = input("→ New time / Enter to skip: ").strip()
            if raw_time in {"/back", "/b"}:
                continue
            if raw_time:
                parsed = parse_update_time(raw_time)
                if parsed is None:
                    print("Please enter a time in HH:MM format, e.g. 09:00.")
                    continue
                update_time = f"{parsed[0]:02d}:{parsed[1]:02d}"

            if not update_timezone:
                update_timezone = _detect_system_timezone()
            print(f"Using this computer's timezone for the schedule: {update_timezone}")
            save_draft_state(update_time=update_time, update_timezone=update_timezone, auto_update=auto_update)
            break

    eta_map = dict(sanitize_feature_eta_dates(feature_eta_dates) or {})
    pace_map = dict(sanitize_expected_tasks_per_week(expected_tasks_per_week) or {})
    features_out = []
    for kw in sanitize_keyword_values(include_keywords):
        entry = {"keyword": kw}
        _, eta_val = _lookup_icase(eta_map, kw)
        if eta_val is not None:
            entry["eta"] = eta_val
        _, pace_val = _lookup_icase(pace_map, kw)
        if pace_val is not None:
            entry["expected_pace"] = pace_val
        features_out.append(entry)
    save_state(
        {
            "features": features_out,
            "exclude": sanitize_keyword_values(excludes),
            "project_keys": [k.strip().upper() for k in project_keys if k.strip()],
            "done_statuses": [s.strip() for s in (done_statuses or DEFAULT_DONE_STATUSES) if s.strip()],
            "output": str(current_output_path()),
            "local_only": local_only,
            "drive_folder": drive_folder,
            "google_client_secrets": make_project_relative(google_client_secrets) if google_client_secrets else None,
            "update_time": update_time,
            "update_timezone": update_timezone,
            "auto_update": auto_update,
        }
    )

    new_features_arg = []
    if edit_run:
        update_run = True
        added_keywords = [k for k in include_keywords if normalize_keyword(k) not in original_keywords_norm]
        new_features_arg = added_keywords

    cmd = [
        sys.executable,
        str(APP_DIR / "jira-report.py"),
        "--state",
        str(STATE_PATH),
        "--output",
        output,
    ]
    if fresh_run:
        cmd.append("--fresh")
    if update_run:
        cmd.append("--update")
    if new_features_arg:
        cmd += ["--new-features", ",".join(new_features_arg)]
    if debug:
        cmd.append("--debug")
    if from_cache:
        cmd.append("--from-cache")
    if filter_pattern:
        cmd += ["--feature-filter", filter_pattern, "--feature-filter-all"]
    for value in excludes:
        cmd.extend(["--exclude", value])

    ensure_jira_token()
    if not _ensure_openpyxl():
        raise SystemExit(
            f"openpyxl is required to build the report. Install it yourself with:\n"
            f"    {sys.executable} -m pip install openpyxl\nthen run this tool again."
        )
    startup_message = "Checking Jira for updates..." if update_run else "Building report..."
    env = os.environ.copy()
    jira_token = read_jira_token()
    if jira_token:
        env["JIRA_TOKEN"] = jira_token
    try:
        returncode, _proc_out = _run_streaming_and_capture(cmd, ROOT, env, startup_message=startup_message)
        result = type("R", (), {"returncode": returncode})()
    except KeyboardInterrupt:
        raise SystemExit(0)
    if update_run:
        vpn_keywords = ("timed out", "connection refused", "Operation timed out",
                        "VPN", "Cannot connect", "network", "dropped")
        if result.returncode == 87 or any(k.lower() in _proc_out.lower() for k in vpn_keywords):
            _send_notification("Can't reach Jira — check your VPN connection.")
        elif result.returncode == 88:
            _send_notification("All quiet on the Jira front. Come back when someone actually does something.")
        elif result.returncode != 0:
            _send_notification("Update finished with errors — check the log.")
        else:
            m = re.search(r"(\d+) tasks? updated their status", _proc_out)
            if m and int(m.group(1)) > 0:
                _send_notification(f"{m.group(1)} task(s) updated their status.")
            else:
                _send_notification("All quiet on the Jira front. Come back when someone actually does something.")
    if result.returncode == 87:
        print_line(
            f"Can't reach Jira — connect to VPN and check that https://{env.get('JIRA_HOST', '')} "
            "opens in your browser, then try again."
        )
        raise SystemExit(result.returncode)
    if result.returncode == 86:
        print_line("Jira sign-in looks stale. Let’s refresh setup once.")
        run_jira_setup()
        jira_token = read_jira_token()
        if jira_token:
            env["JIRA_TOKEN"] = jira_token
        try:
            returncode, _proc_out = _run_streaming_and_capture(cmd, ROOT, env, startup_message=startup_message)
            result = type("R", (), {"returncode": returncode})()
        except KeyboardInterrupt:
            raise SystemExit(0)
    if result.returncode == 87:
        print_line(
            f"Can't reach Jira — connect to VPN and check that https://{env.get('JIRA_HOST', '')} "
            "opens in your browser, then try again."
        )
        raise SystemExit(result.returncode)
    if result.returncode == 88:
        _do_drive_upload(output, local_only, drive_folder, google_client_secrets)
        if auto_update:
            sched_time = update_time or "08:00"
            sched_tz = update_timezone or "UTC"
            status_line(f"✓ Report will be automatically updated daily at {sched_time} {sched_tz}.")
        return
    if result.returncode != 0:
        if "Traceback (most recent call last):" in _proc_out:
            log_path = _write_crash_log(_proc_out)
            print_line(f"Report generation crashed. A log was saved to:\n    {log_path}\nPlease send that file so this can be debugged.")
        else:
            print_line("Report generation stopped. Please adjust the keywords and try again.")
        raise SystemExit(result.returncode)
    status_line("✓ Report built.")

    # Mark today as updated so the missed-update checker doesn't fire
    stamp = APP_DIR / ".last-daily-run-utc"
    stamp.write_text(dt.date.today().isoformat())

    _do_drive_upload(output, local_only, drive_folder, google_client_secrets)

    sched_time = update_time or "08:00"
    sched_tz = update_timezone or "UTC"
    if not auto_update:
        if sys.platform == "darwin" and daily_job_exists():
            subprocess.run(["bash", str(APP_DIR / "uninstall-launchd.sh")], cwd=ROOT, check=False)
        return

    # Check auto-update job health and path
    if sys.platform == "darwin":
        installed_plist = Path.home() / "Library/LaunchAgents" / f"{JOB_NAME}.plist"
        plist_ok = False
        if installed_plist.exists():
            plist_text = installed_plist.read_text()
            plist_ok = str(APP_DIR) in plist_text
        lc = subprocess.run(["launchctl", "list"], capture_output=True, text=True)
        job_line = next((l for l in lc.stdout.splitlines() if JOB_NAME in l), None)
        job_running = job_line and job_line.split()[1] == "0"
        if not plist_ok or not job_running:
            print("Auto-update job needs reinstall — fixing now...")
            subprocess.run(["bash", str(APP_DIR / "install-launchd.sh")], cwd=ROOT, check=False)
            print("✓ Auto-update reinstalled.")

    if daily_job_exists():
        print(f"Next auto-update scheduled at {sched_time} {sched_tz}.")
        ensure_full_disk_access()
    else:
        ensure_full_disk_access()
        run_spinner(
            "Installing daily updates...",
            lambda: subprocess.run(["bash", str(APP_DIR / "install-launchd.sh")], cwd=ROOT, check=True),
        )
        print(f"Auto-update installed. Next run at {sched_time} {sched_tz}.")


_CRASH_LOG_PATH = ROOT / "logs" / "roadmap-crash-log.txt"


def _write_crash_log(body):
    """Shared by both crash paths: an unhandled exception in this process
    itself, and an unhandled crash inside the jira-report.py subprocess it
    launches (see the report-building call in main()) — one file, one
    place to look, regardless of which process actually failed."""
    import platform as _platform
    _CRASH_LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(_CRASH_LOG_PATH, "w") as f:
        f.write(f"Python: {sys.version}\n")
        f.write(f"Platform: {_platform.platform()}\n\n")
        f.write(body)
    return _CRASH_LOG_PATH


def _main_with_crash_log():
    try:
        main()
    except KeyboardInterrupt:
        print()
        raise SystemExit(0)
    except Exception:
        # Anything we didn't already turn into a friendly SystemExit above
        # (those aren't caught here — SystemExit isn't an Exception subclass)
        # is a genuine bug. Save enough to actually debug it, since there's
        # no other log file for the interactive launcher.
        import traceback as _traceback
        log_path = _write_crash_log(_traceback.format_exc())
        print(f"\nSomething went wrong. A log was saved to:\n    {log_path}\nPlease send that file so this can be debugged.")
        raise SystemExit(1)


if __name__ == "__main__":
    _main_with_crash_log()
